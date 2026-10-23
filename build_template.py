"""Builds a clean, generic "Project Status Report" demo template from
scratch: a cover slide, an "In Progress Projects" table, and a "Completed
Projects" table. Deliberately unrelated to any specific industry -- this is
a public-facing demo of the spreadsheet-to-deck workflow, not a real
business template.
"""
import copy

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SLATE = "2D3142"
TEAL = "4F6D7A"
CREAM = "F6F5F3"
LIGHT_GRAY = "E8E6E1"
WHITE = "FFFFFF"
MUTED = "6B7280"

FONT = "Calibri"
OUT_PATH = "templates/base_template.pptx"


def in_(v: float) -> Emu:
    return Emu(int(v * 914400))


def set_bg(slide, hex_color):
    cSld = slide._element.find(qn("p:cSld"))
    existing = cSld.find(qn("p:bg"))
    if existing is not None:
        cSld.remove(existing)
    bg = cSld.makeelement(qn("p:bg"), {})
    bgPr = bg.makeelement(qn("p:bgPr"), {})
    solidFill = bgPr.makeelement(qn("a:solidFill"), {})
    srgbClr = solidFill.makeelement(qn("a:srgbClr"), {"val": hex_color})
    solidFill.append(srgbClr)
    bgPr.append(solidFill)
    bgPr.append(bgPr.makeelement(qn("a:effectLst"), {}))
    bg.append(bgPr)
    cSld.insert(0, bg)


def add_textbox(slide, left, top, width, height, text, size, color, bold=False, align=PP_ALIGN.LEFT, italic=False):
    box = slide.shapes.add_textbox(in_(left), in_(top), in_(width), in_(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = FONT
    run.font.color.rgb = RGBColor.from_string(color)
    return box


def style_cell(cell, text, size=10, bold=False, color="1A1A1A", fill=None, align=PP_ALIGN.LEFT):
    cell.text = text
    cell.margin_left = Emu(45720)
    cell.margin_right = Emu(45720)
    cell.margin_top = Emu(27432)
    cell.margin_bottom = Emu(27432)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(fill)
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    if p.runs:
        run = p.runs[0]
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = FONT
        run.font.color.rgb = RGBColor.from_string(color)


def build_table(slide, left, top, width, headers, template_row_placeholders, col_widths_ratio, extra_blank_rows=10):
    n_cols = len(headers)
    n_rows = 2 + extra_blank_rows  # header + template row + blanks
    row_height = in_(0.4)
    graphic_frame = slide.shapes.add_table(n_rows, n_cols, in_(left), in_(top), in_(width), row_height * n_rows)
    table = graphic_frame.table

    total_ratio = sum(col_widths_ratio)
    for i, ratio in enumerate(col_widths_ratio):
        table.columns[i].width = Emu(int(width * 914400 * (ratio / total_ratio)))

    for c, header in enumerate(headers):
        style_cell(table.cell(0, c), header, size=10, bold=True, color=WHITE, fill=SLATE)

    for c, placeholder in enumerate(template_row_placeholders):
        style_cell(table.cell(1, c), placeholder, size=10, color="1A1A1A", fill=WHITE)

    for r in range(2, n_rows):
        for c in range(n_cols):
            style_cell(table.cell(r, c), "", size=10, fill=WHITE)

    for row in table.rows:
        row.height = row_height

    return graphic_frame


def main():
    prs = Presentation()
    prs.slide_width = in_(10)
    prs.slide_height = in_(7.5)
    blank_layout = prs.slide_layouts[6]

    # ---------- Slide 1: cover ----------
    cover = prs.slides.add_slide(blank_layout)
    set_bg(cover, SLATE)

    band = cover.shapes.add_shape(1, in_(0), in_(0), in_(10), in_(0.12))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor.from_string(TEAL)
    band.line.fill.background()
    band.shadow.inherit = False
    for shp in list(cover.shapes):
        if shp.has_text_frame is False and shp.name.startswith("Rectangle"):
            el = shp._element.find(qn("p:style"))
            if el is not None:
                shp._element.remove(el)

    add_textbox(cover, 0.8, 2.6, 8.4, 0.5, "REPORTING DEMO", 13, TEAL, bold=True)
    add_textbox(cover, 0.8, 3.0, 8.4, 1.0, "{team_name}", 40, WHITE, bold=True)
    add_textbox(cover, 0.8, 3.85, 8.4, 0.6, "Quarterly Project Report", 18, "C9D1D6")

    pill = cover.shapes.add_shape(1, in_(0.8), in_(4.6), in_(2.6), in_(0.42))
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor.from_string(TEAL)
    pill.line.fill.background()
    pill.shadow.inherit = False
    style = pill._element.find(qn("p:style"))
    if style is not None:
        pill._element.remove(style)
    ptf = pill.text_frame
    ptf.word_wrap = True
    ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
    pp = ptf.paragraphs[0]
    pp.alignment = PP_ALIGN.CENTER
    prun = pp.add_run()
    prun.text = "{report_period}"
    prun.font.size = Pt(13)
    prun.font.bold = True
    prun.font.name = FONT
    prun.font.color.rgb = RGBColor.from_string(WHITE)

    # ---------- Slide 2: In Progress Projects ----------
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2, CREAM)
    add_textbox(s2, 0.5, 0.45, 9, 0.6, "In Progress Projects", 26, SLATE, bold=True)
    build_table(
        s2,
        left=0.5,
        top=1.15,
        width=9.0,
        headers=["Project", "Owner", "Start Date", "Target Date", "% Complete", "Notes"],
        template_row_placeholders=[
            "{project_name}",
            "{owner}",
            "{start_date}",
            "{target_date}",
            "{percent_complete}",
            "{notes}",
        ],
        col_widths_ratio=[1.6, 1.0, 1.0, 1.0, 0.9, 1.8],
    )

    # ---------- Slide 3: Completed Projects ----------
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3, CREAM)
    add_textbox(s3, 0.5, 0.45, 9, 0.6, "Completed Projects", 26, SLATE, bold=True)
    build_table(
        s3,
        left=0.5,
        top=1.15,
        width=9.0,
        headers=["Project", "Owner", "Start Date", "Completed Date", "Outcome"],
        template_row_placeholders=[
            "{project_name}",
            "{owner}",
            "{start_date}",
            "{completed_date}",
            "{outcome}",
        ],
        col_widths_ratio=[1.6, 1.0, 1.0, 1.1, 2.6],
    )

    prs.save(OUT_PATH)
    print("saved", OUT_PATH, "slides:", len(prs.slides))


if __name__ == "__main__":
    main()
