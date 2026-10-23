from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT_NAME = "Arial"

HEADER_FILL = PatternFill("solid", fgColor="2D3142")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
REQUIRED_FILL = PatternFill("solid", fgColor="FEF3C7")
EXAMPLE_FILL = PatternFill("solid", fgColor="F3F4F6")
EXAMPLE_FONT = Font(name=FONT_NAME, italic=True, color="6B7280", size=10)
BODY_FONT = Font(name=FONT_NAME, size=10)
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# (header, key, width, group, required, number_format)
COLUMNS = [
    ("Team Name", "team_name", 20, "all", True, None),
    ("Status", "status", 14, "all", True, None),
    ("Project Name", "project_name", 26, "all", True, None),
    ("Owner", "owner", 16, "all", False, None),
    ("Start Date", "start_date", 14, "all", False, "yyyy-mm-dd"),
    ("Target Date", "target_date", 14, "in_progress", False, "yyyy-mm-dd"),
    ("Percent Complete", "percent_complete", 16, "in_progress", False, "0%"),
    ("Notes", "notes", 36, "in_progress", False, None),
    ("Completed Date", "completed_date", 16, "completed", False, "yyyy-mm-dd"),
    ("Outcome", "outcome", 40, "completed", False, None),
]

EXAMPLE_ROWS = [
    {
        "team_name": "Marketing Team",
        "status": "in_progress",
        "project_name": "Website Redesign",
        "owner": "Alex Kim",
        "start_date": "2026-01-10",
        "target_date": "2026-05-01",
        "percent_complete": 0.65,
        "notes": "On track; new homepage mockups approved.",
    },
    {
        "team_name": "Marketing Team",
        "status": "completed",
        "project_name": "Customer Survey 2025",
        "owner": "Priya Patel",
        "start_date": "2025-09-01",
        "completed_date": "2025-11-15",
        "outcome": "Collected 1,200 responses; results shared with leadership.",
    },
    {
        "team_name": "Engineering Team",
        "status": "in_progress",
        "project_name": "Mobile App Launch",
        "owner": "Sam Rivera",
        "start_date": "2026-01-15",
        "target_date": "2026-06-01",
        "percent_complete": 0.3,
        "notes": "Beta testing begins next sprint.",
    },
]

wb = Workbook()

# ---------- Sheet 1: Instructions ----------
ws = wb.active
ws.title = "Instructions"
ws.sheet_view.showGridLines = False

ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 12
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 46
ws.column_dimensions["E"].width = 34

ws["A1"] = "Reporting Demo — Spreadsheet Template"
ws["A1"].font = Font(name=FONT_NAME, bold=True, size=14)
ws.merge_cells("A1:E1")

ws["A2"] = (
    "Fill in the “Projects” tab — one row per team per project — "
    "then upload this file on the demo site. It will produce one personalized "
    "PowerPoint/PDF report per team."
)
ws["A2"].font = Font(name=FONT_NAME, italic=True, color="374151", size=10)
ws.merge_cells("A2:E2")
ws.row_dimensions[2].height = 30
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")

ws["A4"] = "How to use this file"
ws["A4"].font = Font(name=FONT_NAME, bold=True, size=11)
steps = [
    f"1. Go to the “Projects” tab.",
    f"2. Delete the {len(EXAMPLE_ROWS)} example rows (shaded gray) — they're just there to show the expected format.",
    "3. Add one row per team per project. The same team can appear on multiple rows "
    "(one per project).",
    "4. “Status” must be exactly “in_progress” or “completed” — use the dropdown in that column.",
    "5. Leave a cell blank if it doesn't apply — don't delete columns.",
    "6. Save the file and upload it on the site.",
]
for i, s in enumerate(steps):
    cell = ws.cell(row=5 + i, column=1, value=s)
    cell.font = Font(name=FONT_NAME, size=10)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=5 + i, start_column=1, end_row=5 + i, end_column=5)
    ws.row_dimensions[5 + i].height = 16

legend_row = 5 + len(steps) + 1
ws.cell(row=legend_row, column=1, value="Column reference").font = Font(name=FONT_NAME, bold=True, size=11)

header_row = legend_row + 1
headers = ["Column", "Required?", "Used for", "Format / example", "Notes"]
for c, h in enumerate(headers, start=1):
    cell = ws.cell(row=header_row, column=c, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.border = BORDER
    cell.alignment = Alignment(vertical="center")
ws.row_dimensions[header_row].height = 18

notes_by_key = {
    "team_name": "Exact name to print on that team's cover slide.",
    "status": "Controls which table (in progress vs. completed) the row lands in.",
    "project_name": "Project name.",
    "percent_complete": "Enter as 0.65 or 65 or “65%” — all work.",
    "notes": "Free text, in-progress projects only.",
    "outcome": "Free text, completed projects only.",
}
format_examples = {
    "start_date": "2026-01-10",
    "target_date": "2026-05-01",
    "completed_date": "2025-11-15",
    "percent_complete": "0.65 → 65%",
    "status": "in_progress | completed",
}

r = header_row + 1
for header, key, _w, group, required, _fmt in COLUMNS:
    ws.cell(row=r, column=1, value=header).font = BODY_FONT
    req_cell = ws.cell(row=r, column=2, value="Required" if required else "Optional")
    req_cell.font = Font(name=FONT_NAME, size=10, bold=required, color="B45309" if required else "6B7280")
    ws.cell(row=r, column=3, value=group).font = BODY_FONT
    ws.cell(row=r, column=4, value=format_examples.get(key, "")).font = BODY_FONT
    ws.cell(row=r, column=5, value=notes_by_key.get(key, "")).font = BODY_FONT
    for c in range(1, 6):
        cell = ws.cell(row=r, column=c)
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    if required:
        for c in range(1, 6):
            ws.cell(row=r, column=c).fill = REQUIRED_FILL
    r += 1

# ---------- Sheet 2: Projects (the data-entry tab) ----------
ws2 = wb.create_sheet("Projects")
ws2.sheet_view.showGridLines = False
ws2.freeze_panes = "A2"

for c, (header, key, width, _group, required, _fmt) in enumerate(COLUMNS, start=1):
    letter = get_column_letter(c)
    ws2.column_dimensions[letter].width = width
    cell = ws2.cell(row=1, column=c, value=header)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.border = BORDER
    cell.alignment = Alignment(vertical="center", wrap_text=True)

ws2["A1"].comment = Comment(
    f"One row per team per project. See the Instructions tab for the full column "
    f"reference. Delete the {len(EXAMPLE_ROWS)} gray example rows before uploading your real data.",
    "Reporting Demo",
    width=260,
    height=100,
)

for row_i, example in enumerate(EXAMPLE_ROWS, start=2):
    for c, (_header, key, _w, _group, _required, number_format) in enumerate(COLUMNS, start=1):
        value = example.get(key)
        cell = ws2.cell(row=row_i, column=c, value=value)
        cell.font = EXAMPLE_FONT
        cell.fill = EXAMPLE_FILL
        cell.border = BORDER
        if number_format:
            cell.number_format = number_format

first_blank = len(EXAMPLE_ROWS) + 2
last_blank = first_blank + 300
for row_i in range(first_blank, last_blank):
    for c, (_header, _key, _w, _group, _required, number_format) in enumerate(COLUMNS, start=1):
        cell = ws2.cell(row=row_i, column=c)
        cell.font = BODY_FONT
        cell.border = BORDER
        if number_format:
            cell.number_format = number_format

status_col_letter = get_column_letter([k for _, k, *_ in COLUMNS].index("status") + 1)
dv = DataValidation(type="list", formula1='"in_progress,completed"', allow_blank=True, showDropDown=False)
dv.error = "Status must be either in_progress or completed."
dv.errorTitle = "Invalid status"
ws2.add_data_validation(dv)
dv.add(f"{status_col_letter}2:{status_col_letter}{last_blank - 1}")

out_path = "/tmp/claude-0/-home-user-wedding-site/d3461d81-29aa-508f-b6fc-0344a8e10dce/scratchpad/reporting-demo/report_template.xlsx"
wb.save(out_path)
print("saved", out_path)
