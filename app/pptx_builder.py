"""Fills the base_template.pptx merge fields for a single team.

Slides are found by content, not a fixed position:
    cover slide            - has a shape whose text is exactly "{team_name}"
    in-progress table      - has a table whose header row contains "% Complete"
    completed table        - has a table whose header row contains "Outcome"
"""
from __future__ import annotations

import copy
import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

LAST_ROW_SHADE = RGBColor(0xE8, 0xE6, 0xE1)
COVER_MARKER = "{team_name}"
IN_PROGRESS_TABLE_MARKER = "% Complete"
COMPLETED_TABLE_MARKER = "Outcome"


def _replace_placeholders(text_frame, mapping: dict) -> None:
    for paragraph in text_frame.paragraphs:
        runs = paragraph.runs
        if not runs:
            continue
        full_text = "".join(r.text for r in runs)
        if "{" not in full_text:
            continue
        new_text = full_text
        for key, value in mapping.items():
            new_text = new_text.replace("{" + key + "}", value)
        if new_text == full_text:
            continue
        runs[0].text = new_text
        for r in runs[1:]:
            r.text = ""


def _slide_has_exact_text(slide, text: str) -> bool:
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip() == text:
            return True
    return False


def _slide_table_header_contains(slide, needle: str) -> bool:
    for shape in slide.shapes:
        if shape.has_table:
            header_cells = [c.text for c in shape.table.rows[0].cells]
            if any(needle in cell for cell in header_cells):
                return True
    return False


def _find_slide(prs, predicate, description: str):
    for slide in prs.slides:
        if predicate(slide):
            return slide
    raise ValueError(f"Could not find the {description} in this template.")


def _find_table_shape(slide):
    for shape in slide.shapes:
        if shape.has_table:
            return shape
    raise ValueError("Expected slide to contain a table, none found.")


def _shade_and_bold_row(row) -> None:
    for cell in row.cells:
        cell.fill.solid()
        cell.fill.fore_color.rgb = LAST_ROW_SHADE
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True


def _fill_table(slide, rows: list[dict]) -> None:
    table_shape = _find_table_shape(slide)
    table = table_shape.table
    tbl_elem = table._tbl
    trs = tbl_elem.findall(qn("a:tr"))
    if len(trs) < 2:
        raise ValueError("Template table must have a header row and a merge-field row.")

    template_tr = trs[1]
    for extra_tr in trs[2:]:
        tbl_elem.remove(extra_tr)
    tbl_elem.remove(template_tr)

    for _ in rows:
        new_tr = copy.deepcopy(template_tr)
        tbl_elem.append(new_tr)

    data_rows = list(table.rows)[1:]
    for row_data, row in zip(rows, data_rows):
        for cell in row.cells:
            _replace_placeholders(cell.text_frame, row_data)

    if len(data_rows) > 1:
        _shade_and_bold_row(data_rows[-1])


def build_team_deck(
    template_path: str,
    team_name: str,
    report_period: str,
    in_progress_rows: list[dict],
    completed_rows: list[dict],
) -> io.BytesIO:
    prs = Presentation(template_path)

    cover_slide = _find_slide(
        prs, lambda s: _slide_has_exact_text(s, COVER_MARKER), "cover slide"
    )
    for shape in cover_slide.shapes:
        if shape.has_text_frame:
            _replace_placeholders(
                shape.text_frame, {"team_name": team_name, "report_period": report_period}
            )

    in_progress_slide = _find_slide(
        prs,
        lambda s: _slide_table_header_contains(s, IN_PROGRESS_TABLE_MARKER),
        "In Progress Projects table slide",
    )
    completed_slide = _find_slide(
        prs,
        lambda s: _slide_table_header_contains(s, COMPLETED_TABLE_MARKER),
        "Completed Projects table slide",
    )

    _fill_table(in_progress_slide, in_progress_rows)
    _fill_table(completed_slide, completed_rows)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf
